"""EP1 ACH2118 - stacking para classificar a clareza de respostas do e-SIC {c1, c234, c5}.

Script unico e independente. Coloque train.xlsx e test1.xlsx nesta pasta (ou na pasta de cima) e rode:

    python stacking_ep1.py                    # versao MiniLM (melhor resultado)
    python stacking_ep1.py --versao e5        # versao e5-small
    python stacking_ep1.py --versao ambas     # as duas, uma depois da outra

O que ele faz:
    1. mede o baseline oficial e o stacking em validacao cruzada de 5 dobras (semente 123);
    2. mostra a acuracia em cada dobra;
    3. treina com todo o treino e grava a planilha de teste rotulada em saida/test1_<versao>.xlsx.

Modelo: quatro modelos-base + meta-classificador (regressao logistica).
    lr_pal    regressao logistica sobre TF-IDF de palavras (1-2-gramas)
    svc       SVM linear sobre TF-IDF de palavras (1-2) + caracteres (2-5)
    emb       regressao logistica sobre embeddings de sentenca (MiniLM ou e5-small, congelado)
    extremos  regressao logistica c1 x c5 treinada so nos extremos (eixo "ruim-bom")
"""
import argparse
import re
import time
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import scipy.sparse as sp
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.model_selection import StratifiedKFold
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.svm import LinearSVC

warnings.filterwarnings("ignore")
SEED = 123
N_DOBRAS = 5
ROTULOS = ["c1", "c234", "c5"]
AQUI = Path(__file__).resolve().parent
# versao -> (modelo no Hugging Face, prefixo exigido pelo modelo, maximo de subpalavras)
EMBEDDINGS = {
    "minilm": ("sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2", "", 256),
    "e5": ("intfloat/multilingual-e5-small", "query: ", 256),
}


# ----------------------------------------------------------------------------- dados
def achar(nome):
    for pasta in (AQUI, AQUI.parent):
        if (pasta / nome).exists():
            return pasta / nome
    raise FileNotFoundError(f"{nome} nao encontrado em {AQUI} nem em {AQUI.parent}")


def limpar(t):
    """Unico pre-processamento: junta espacos e quebras de linha repetidos."""
    return re.sub(r"\s+", " ", str(t)).strip()


def carregar():
    tr, te = pd.read_excel(achar("train.xlsx")), pd.read_excel(achar("test1.xlsx"))
    bruto = np.array([str(t) for t in tr["resp_text"]] + [str(t) for t in te["resp_text"]])
    texto = np.array([limpar(t) for t in bruto])
    y = tr["clarity"].map({r: i for i, r in enumerate(ROTULOS)}).values.astype(int)
    return bruto, texto, y


# ----------------------------------------------------------------------------- representacoes
def tfidf(texto):
    """TF-IDF de palavras e de caracteres, ajustado em treino + teste (nao usa rotulos)."""
    pal = TfidfVectorizer(ngram_range=(1, 2), min_df=2, sublinear_tf=True).fit_transform(texto)
    car = TfidfVectorizer(analyzer="char_wb", ngram_range=(2, 5), min_df=3, sublinear_tf=True).fit_transform(texto)
    return pal.tocsr().astype(np.float32), sp.hstack([pal, car]).tocsr().astype(np.float32)


def embeddings(versao, texto):
    """Vetores de sentenca do modelo pequeno (CPU). Ficam em cache/emb_<versao>.npy para nao recalcular."""
    arq = AQUI / "cache" / f"emb_{versao}.npy"
    if arq.exists():
        E = np.load(arq)
        if len(E) == len(texto):
            print(f"embeddings '{versao}' lidos do cache", flush=True)
            return E
    import torch
    from sentence_transformers import SentenceTransformer
    nome, prefixo, max_len = EMBEDDINGS[versao]
    print(f"extraindo embeddings com {nome} (20 a 30 min em CPU; so na primeira vez)...", flush=True)
    torch.set_num_threads(4)
    modelo = SentenceTransformer(nome, device="cpu")
    modelo.max_seq_length = max_len
    entrada, partes, t0 = [prefixo + t for t in texto], [], time.time()
    for i in range(0, len(entrada), 2000):
        partes.append(modelo.encode(entrada[i:i + 2000], batch_size=16, show_progress_bar=False,
                                    normalize_embeddings=True, convert_to_numpy=True))
        print(f"  {min(i + 2000, len(entrada))}/{len(entrada)} textos em {time.time() - t0:.0f}s", flush=True)
    E = np.vstack(partes).astype(np.float32)
    arq.parent.mkdir(exist_ok=True)
    np.save(arq, E)
    return E


# ----------------------------------------------------------------------------- modelos-base
# Cada funcao treina nas linhas `a` (rotulos ya) e devolve as pontuacoes das linhas `b`.
def base_lr_pal(X, a, ya, b):
    return LogisticRegression(C=1.0, max_iter=300).fit(X["pal"][a], ya).predict_log_proba(X["pal"][b])


def base_svc(X, a, ya, b):
    return LinearSVC(C=0.1, random_state=SEED).fit(X["tfidf"][a], ya).decision_function(X["tfidf"][b])


def base_emb(X, a, ya, b):
    m = make_pipeline(StandardScaler(), LogisticRegression(C=0.01, max_iter=300)).fit(X["emb"][a], ya)
    return np.log(m.predict_proba(X["emb"][b]) + 1e-6)


def base_extremos(X, a, ya, b):
    ext = ya != 1                                   # so c1 e c5
    m = LogisticRegression(C=1.0, max_iter=300).fit(X["tfidf"][a[ext]], ya[ext])
    return m.decision_function(X["tfidf"][b]).reshape(-1, 1)


BASES = {"lr_pal": base_lr_pal, "svc": base_svc, "emb": base_emb, "extremos": base_extremos}


def meta():
    return make_pipeline(StandardScaler(), LogisticRegression(C=0.1, max_iter=3000))


def baseline_oficial(bruto, y, a, b):
    """Baseline do professor: TF-IDF padrao + regressao logistica balanceada, max_iter=100."""
    v = TfidfVectorizer()
    m = LogisticRegression(class_weight="balanced", max_iter=100, random_state=SEED).fit(v.fit_transform(bruto[a]), y[a])
    return m.predict(v.transform(bruto[b]))


# ----------------------------------------------------------------------------- execucao
def rodar(versao, bruto, texto, y, pal, cheio, cache_bases):
    n = len(y)
    treino, teste = np.arange(n), np.arange(n, len(texto))
    X = {"pal": pal, "tfidf": cheio, "emb": embeddings(versao, texto)}
    dobra = np.zeros(n, dtype=int)
    for k, (_, v) in enumerate(StratifiedKFold(N_DOBRAS, shuffle=True, random_state=SEED).split(y, y)):
        dobra[v] = k

    # 1) previsoes "fora da dobra" de cada modelo-base (treina em 4 dobras, preve a quinta) e do teste
    fora, no_teste = {}, {}
    for nome, f in BASES.items():
        if nome != "emb" and nome in cache_bases:    # os blocos de TF-IDF sao iguais nas duas versoes
            fora[nome], no_teste[nome] = cache_bases[nome]
            continue
        t0 = time.time()
        partes = [f(X, treino[dobra != k], y[dobra != k], treino[dobra == k]) for k in range(N_DOBRAS)]
        o = np.zeros((n, partes[0].shape[1]))
        for k in range(N_DOBRAS):
            o[dobra == k] = partes[k]
        fora[nome], no_teste[nome] = o, f(X, treino, y, teste)
        if nome != "emb":
            cache_bases[nome] = (fora[nome], no_teste[nome])
        print(f"  modelo-base {nome} pronto ({time.time() - t0:.0f}s)", flush=True)
    Xm = np.hstack([fora[b] for b in BASES])
    Xm_teste = np.hstack([no_teste[b] for b in BASES])

    # 2) validacao cruzada do meta-classificador e do baseline oficial, nas mesmas dobras
    if "baseline" not in cache_bases:
        pb = np.zeros(n, dtype=int)
        for k in range(N_DOBRAS):
            pb[dobra == k] = baseline_oficial(bruto, y, treino[dobra != k], treino[dobra == k])
        cache_bases["baseline"] = pb
    pb, pm = cache_bases["baseline"], np.zeros(n, dtype=int)
    for k in range(N_DOBRAS):
        pm[dobra == k] = meta().fit(Xm[dobra != k], y[dobra != k]).predict(Xm[dobra == k])
    tab = pd.DataFrame({"dobra": [str(k + 1) for k in range(N_DOBRAS)],
                        "baseline_oficial": [(pb[dobra == k] == y[dobra == k]).mean() for k in range(N_DOBRAS)],
                        "stacking": [(pm[dobra == k] == y[dobra == k]).mean() for k in range(N_DOBRAS)]})
    tab.loc[len(tab)] = ["media", tab.baseline_oficial.mean(), tab.stacking.mean()]
    tab["diferenca"] = tab.stacking - tab.baseline_oficial
    print(f"\n=== versao {versao}: acuracia em validacao cruzada de {N_DOBRAS} dobras (semente {SEED}) ===")
    print(tab.assign(**{c: (tab[c] * 100).round(2) for c in tab.columns[1:]}).to_string(index=False))

    # 3) modelo final: meta-classificador treinado em todas as previsoes fora da dobra -> rotula o teste
    pred = meta().fit(Xm, y).predict(Xm_teste)
    (AQUI / "saida").mkdir(exist_ok=True)
    tab.to_csv(AQUI / "saida" / f"resultado_{versao}.csv", index=False)
    from openpyxl import load_workbook
    wb = load_workbook(achar("test1.xlsx"))
    ws = wb.active
    assert ws.max_row - 1 == len(pred), "numero de linhas diferente do teste"
    for i, p in enumerate(pred):
        ws.cell(row=i + 2, column=2, value=ROTULOS[int(p)])
    saida = AQUI / "saida" / f"test1_{versao}.xlsx"
    wb.save(saida)
    print(f"teste rotulado em {saida} | distribuicao:",
          {ROTULOS[i]: int(q) for i, q in enumerate(np.bincount(pred, minlength=3))}, "\n", flush=True)


def main():
    ap = argparse.ArgumentParser(description="Stacking do EP1 (ACH2118)")
    ap.add_argument("--versao", choices=["minilm", "e5", "ambas"], default="minilm")
    versoes = ["minilm", "e5"] if ap.parse_args().versao == "ambas" else [ap.parse_args().versao]
    t0 = time.time()
    bruto, texto, y = carregar()
    print(f"{len(y)} respostas de treino, {len(texto) - len(y)} de teste. Calculando TF-IDF (alguns minutos)...", flush=True)
    pal, cheio = tfidf(texto)
    cache_bases = {}
    for v in versoes:
        rodar(v, bruto, texto, y, pal, cheio, cache_bases)
    print(f"tempo total: {(time.time() - t0) / 60:.1f} min")


if __name__ == "__main__":
    main()
