"""Funcoes compartilhadas por todas as abordagens do EP1 (ACH2118).

Estrutura esperada (pasta experimentos/):
    train.xlsx, test1.xlsx          -> dados do professor (copiar para dentro de experimentos/)
    codigo/                         -> scripts
    dados/folds.csv                 -> divisao fixa da validacao cruzada
    resultados/, predicoes/         -> saidas
"""
from pathlib import Path
import re
import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
LABELS = ["c1", "c234", "c5"]
N_FOLDS = 5
SEED = 42


def clean(t):
    """Pre-processamento minimo: junta espacos/quebras de linha repetidos."""
    return re.sub(r"\s+", " ", str(t)).strip()


def load_data():
    """Devolve (train, test). train tem colunas text, y (0,1,2) e fold (0..4)."""
    (ROOT / "dados").mkdir(exist_ok=True)
    tr = pd.read_excel(ROOT / "train.xlsx")
    te = pd.read_excel(ROOT / "test1.xlsx")
    tr = pd.DataFrame({"text": tr["resp_text"].map(clean),
                       "y": tr["clarity"].map({l: i for i, l in enumerate(LABELS)}).astype(int)})
    te = pd.DataFrame({"text": te["resp_text"].map(clean)})

    fpath = ROOT / "dados" / "folds.csv"
    if fpath.exists():
        tr["fold"] = pd.read_csv(fpath)["fold"].values
    else:
        # Textos identicos ficam sempre na mesma dobra (grupo = texto), para que
        # uma resposta repetida nunca apareca no treino e na validacao ao mesmo tempo.
        from sklearn.model_selection import StratifiedGroupKFold
        groups = pd.factorize(tr["text"])[0]
        cv = StratifiedGroupKFold(n_splits=N_FOLDS, shuffle=True, random_state=SEED)
        fold = np.zeros(len(tr), dtype=int)
        for k, (_, va) in enumerate(cv.split(tr["text"], tr["y"], groups)):
            fold[va] = k
        tr["fold"] = fold
        pd.DataFrame({"linha": np.arange(len(tr)), "fold": fold}).to_csv(fpath, index=False)
    return tr, te


def save_test_xlsx(pred_idx, out_name):
    """Grava uma copia do test1.xlsx com a coluna clarity preenchida (formato original)."""
    from openpyxl import load_workbook
    wb = load_workbook(ROOT / "test1.xlsx")
    ws = wb.active
    assert ws.max_row - 1 == len(pred_idx), "numero de linhas diferente do teste"
    for i, p in enumerate(pred_idx):
        ws.cell(row=i + 2, column=2, value=LABELS[int(p)])
    (ROOT / "predicoes").mkdir(exist_ok=True)
    out = ROOT / "predicoes" / out_name
    wb.save(out)
    return out
